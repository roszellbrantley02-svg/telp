"""
lattice/documents.py - reading the user's own files into the one memory.

Drop a PDF, a Word file or a plain text file on Telp and he reads it the
way he reads a web page: the prose is split into passages, each passage
becomes a memory row tagged `doc:<file name>` (with its page number in
the row's metadata), and the claim store reads the same rows so the
fact routes can answer from them. Questions then go through the normal
retrieval path, and "how do you know?" cites the file.

Reading the same file again replaces what was read before, so an edited
file never leaves stale passages behind. `forget_document` erases a
file on command. Nothing here needs the network.

    read_document(agent, "C:/Users/me/report.pdf")
    -> {"title": "report.pdf", "source": "doc:report.pdf", "pages": 12,
        "added": 143, "replaced": 0, "error": None}
"""
from __future__ import annotations

import re
from pathlib import Path

SUPPORTED = (".pdf", ".docx", ".txt", ".md")
SOURCE_PREFIX = "doc:"

# a passage is a memory row: long enough to carry a claim, short enough
# to be one thought (the same band the web and video readers use)
MIN_PASSAGE, MAX_PASSAGE = 30, 320


class DocumentError(Exception):
    """A file Telp cannot read, with a reason the user can act on."""


# ─── extraction ─────────────────────────────────────────────────────

def extract_pages(path: str | Path) -> list[tuple[int | None, str]]:
    """The file's text as (page number, text) pairs. Word and text files
    have no pages and come back as one (None, text) pair."""
    p = Path(path)
    if not p.exists():
        raise DocumentError(f"there's no file at {p}")
    if p.is_dir():
        raise DocumentError(f"{p.name} is a folder, not a file")
    ext = p.suffix.lower()
    if ext not in SUPPORTED:
        raise DocumentError(
            f"I can't read {ext or 'that kind of'} files yet - I read "
            + ", ".join(SUPPORTED))
    if ext == ".pdf":
        return _pdf_pages(p)
    if ext == ".docx":
        return [(None, _docx_text(p))]
    return [(None, p.read_text(encoding="utf-8", errors="replace"))]


def _pdf_pages(p: Path) -> list[tuple[int | None, str]]:
    try:
        from pypdf import PdfReader
    except ImportError as e:                    # pragma: no cover
        raise DocumentError("PDF reading needs the pypdf package "
                            "(pip install pypdf)") from e
    try:
        reader = PdfReader(str(p))
        if reader.is_encrypted:
            # an empty password opens most "protected" PDFs; others
            # really are locked
            try:
                reader.decrypt("")
            except Exception:
                raise DocumentError(f"{p.name} is password-protected")
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            pages.append((i, text))
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError(f"{p.name} doesn't open as a PDF ({e})") from e
    if not any(t.strip() for _, t in pages):
        raise DocumentError(
            f"{p.name} has no readable text - it may be a scanned image")
    return pages


def _docx_text(p: Path) -> str:
    try:
        import docx
    except ImportError as e:                    # pragma: no cover
        raise DocumentError("Word reading needs the python-docx package "
                            "(pip install python-docx)") from e
    try:
        d = docx.Document(str(p))
    except Exception as e:
        raise DocumentError(f"{p.name} doesn't open as a Word file ({e})") \
            from e
    parts = [para.text for para in d.paragraphs]
    # tables read row by row, cells joined, each row its own paragraph
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            cells = list(dict.fromkeys(c for c in cells if c))
            if cells:
                parts.append("")
                parts.append(" - ".join(cells))
    text = "\n".join(parts)
    if not text.strip():
        raise DocumentError(f"{p.name} has no text in it")
    return text


# ─── passages ───────────────────────────────────────────────────────

_SENTENCE_END = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9\"'(\[])")
_LEANS_ON_PREVIOUS = re.compile(
    r"^(?:It|Its|They|Their|This|These|That|Those|He|She|His|Her|Here|"
    r"There|Then|So|Also|However|Instead|Otherwise)\b")


def split_passages(text: str) -> list[str]:
    """Prose passages of MIN..MAX characters: sentences, with short ones
    joined to their neighbours and fragments (headings, page furniture,
    tables of numbers) left out."""
    out: list[str] = []
    for block in _blocks(text):
        # one sentence per row keeps retrieval sharp and lets the claim
        # reader work; a short sentence, or one that leans on the last
        # ("It cost 4.2 million dollars"), rides along with it instead
        buf = ""
        for s in _SENTENCE_END.split(block):
            s = s.strip()
            if not s:
                continue
            leans = _LEANS_ON_PREVIOUS.match(s) is not None
            if buf and (len(s) < MIN_PASSAGE + 10 or leans) \
                    and len(buf) + len(s) < MAX_PASSAGE:
                buf = buf + " " + s
                continue
            _emit(out, buf)
            buf = s
        _emit(out, buf)
    return out


_HEADING_WORDS = 8


def _blocks(text: str) -> list[str]:
    """Paragraphs. A blank line ends one; so does a heading (a short line
    with no closing punctuation, followed by a new sentence) or a bullet.
    PDF line breaks fall inside sentences, so other lines are joined and
    hyphenated line ends ("infor-\\nmation") become one word."""
    # ("infor-\nmation") -> "information"
    text = re.sub(r"-\n(?=[a-z])", "", text)
    lines = [ln.strip() for ln in text.split("\n")]
    blocks: list[str] = []
    cur: list[str] = []

    def flush():
        if cur:
            blocks.append(" ".join(cur))
            cur.clear()

    for i, ln in enumerate(lines):
        if not ln:
            flush()
            continue
        if re.match(r"[-*\u2022]\s", ln):
            flush()
        cur.append(ln)
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        heading = (len(ln.split()) <= _HEADING_WORDS
                   and not re.search(r"[.!?,;:]$", ln)
                   and (not nxt or nxt[:1].isupper() or nxt[:1].isdigit()))
        if heading:
            flush()
    flush()
    return [re.sub(r"\s+", " ", b).strip() for b in blocks]


def _emit(out: list[str], passage: str) -> None:
    passage = passage.strip()
    if not (MIN_PASSAGE <= len(passage) <= MAX_PASSAGE):
        if len(passage) > MAX_PASSAGE:
            # one very long sentence: keep its head rather than lose it
            passage = passage[:MAX_PASSAGE].rsplit(" ", 1)[0] + " ..."
        else:
            return
    letters = sum(c.isalpha() for c in passage)
    if letters < len(passage) * 0.55:           # numbers, dots, rules
        return
    if not re.search(r"[a-z]{3,}\s+[a-z]{2,}", passage):  # words, not codes
        return
    out.append(passage)


# ─── memory ─────────────────────────────────────────────────────────

def source_for(path: str | Path) -> str:
    return SOURCE_PREFIX + Path(path).name


def documents(agent) -> list[dict]:
    """The files Telp has read: [{name, passages}], oldest first."""
    seen: dict[str, int] = {}
    for src in agent.lattice._sources:
        if src.startswith(SOURCE_PREFIX):
            name = src[len(SOURCE_PREFIX):]
            seen[name] = seen.get(name, 0) + 1
    return [{"name": n, "passages": c} for n, c in seen.items()]


def find_document(agent, name: str) -> str | None:
    """The read file whose name matches `name` (exact, then without the
    extension, then as a substring), or None."""
    name = name.strip().strip("'\"").lower()
    if not name:
        return None
    names = [d["name"] for d in documents(agent)]
    for n in names:
        if n.lower() == name:
            return n
    for n in names:
        if Path(n).stem.lower() in (name, Path(name).stem):
            return n
    hits = [n for n in names if name in n.lower()]
    return hits[0] if len(hits) == 1 else None


def forget_document(agent, name: str) -> list[str]:
    """Erase every passage read from the file. Returns the texts gone."""
    src = SOURCE_PREFIX + name
    ids = [mid for mid, s in zip(agent.lattice._ids, agent.lattice._sources)
           if s == src]
    gone = agent.lattice.delete_ids(ids)
    _rebuild_claims(agent)
    return gone


def read_document(agent, path: str | Path, max_passages: int = 600) -> dict:
    """Read a file into memory. Re-reading replaces the earlier reading."""
    p = Path(path).expanduser()
    title = p.name
    try:
        pages = extract_pages(p)
    except DocumentError as e:
        return {"title": title, "source": source_for(p), "pages": 0,
                "added": 0, "replaced": 0, "error": str(e)}
    src = source_for(p)
    replaced = len(forget_document(agent, p.name)) if any(
        s == src for s in agent.lattice._sources) else 0
    n = 0
    seen: set[str] = set()
    for page_no, text in pages:
        for passage in split_passages(text):
            if passage in seen:
                continue
            seen.add(passage)
            meta = {"page": page_no} if page_no else {}
            agent.lattice.add(passage, source=src, **meta)
            try:
                agent.structured.add_sentence(passage, source=src)
            except Exception:
                pass
            n += 1
            if n >= max_passages:
                break
        if n >= max_passages:
            break
    page_count = len(pages) if pages and pages[0][0] is not None else 0
    err = None if n else "no readable prose in it"
    return {"title": title, "source": src, "pages": page_count, "added": n,
            "replaced": replaced, "error": err}


def _rebuild_claims(agent) -> None:
    try:
        agent._rebuild_structured_qa()
    except Exception:
        pass
