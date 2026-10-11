"""Reading the user's own files: PDF, Word and text, into the one memory."""
import subprocess
import sys
from pathlib import Path

import pytest

from tests.test_cli import _telp


# ─── sample files (built here: no binary fixtures in the repo) ──────

def make_docx(path: Path, paragraphs: list[str], table: list[list[str]] = ()):
    import docx
    d = docx.Document()
    d.add_heading("Quarterly report", 1)
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    d.save(str(path))


def make_pdf(path: Path, pages: list[str]) -> None:
    """A minimal text PDF (Helvetica, one Tj per line) pypdf can read."""
    objs: list[str] = []

    def add(o: str) -> int:
        objs.append(o)
        return len(objs)

    font = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    kids = []
    for text in pages:
        lines = text.split("\n")
        stream = ("BT /F1 12 Tf 50 750 Td 14 TL " + " ".join(
            "(" + ln.replace("\\", "\\\\").replace("(", "\\(")
            .replace(")", "\\)") + ") Tj T*" for ln in lines) + " ET")
        c = add(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        kids.append(add(f"<< /Type /Page /Parent PARENT /MediaBox [0 0 612 792] "
                        f"/Contents {c} 0 R /Resources << /Font << /F1 {font} 0 R "
                        f">> >> >>"))
    pages_obj = add(f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] "
                    f"/Count {len(kids)} >>")
    objs = [o.replace("PARENT", f"{pages_obj} 0 R") for o in objs]
    cat = add(f"<< /Type /Catalog /Pages {pages_obj} 0 R >>")
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root {cat} 0 R >>\n"
            f"startxref\n{x}\n%%EOF\n").encode()
    path.write_bytes(out)


REPORT = [
    "The Harlow bridge project finished in March 2024. It cost 4.2 million "
    "dollars and was led by Dana Okafor.",
    "Next year the team plans to repaint the Harlow bridge. The budget for "
    "that is 300 thousand dollars.",
]
MANUAL = [
    "Page 1\nThe Vexlo heater must be switched off before cleaning. Use a\n"
    "damp cloth and never use solvents on the casing.",
    "Page 2\nThe warranty on the Vexlo heater lasts three years from the "
    "date of\npurchase. Keep your receipt as proof.",
]


@pytest.fixture
def files(tmp_path):
    d = tmp_path / "files"
    d.mkdir()
    make_docx(d / "report.docx", REPORT,
              [["Item", "Cost"], ["Steel girders", "1.1 million dollars"]])
    make_pdf(d / "manual.pdf", MANUAL)
    (d / "notes.txt").write_text("Rehearsal is on Thursday at the old hall.\n"
                                 "Bring the blue folder with the sheet music.\n")
    return d


# ─── extraction and passages ────────────────────────────────────────

def test_word_file_reads_paragraphs_and_table_rows_headings_dropped(files):
    from lattice.documents import extract_pages, split_passages
    (page, text), = extract_pages(files / "report.docx")
    assert page is None
    passages = split_passages(text)
    assert passages[0].startswith("The Harlow bridge project finished")
    assert "Dana Okafor" in passages[0]            # "It cost..." rides along
    assert "Steel girders - 1.1 million dollars" in passages
    assert not any(p.startswith("Quarterly report") for p in passages)
    assert not any("Item - Cost" in p for p in passages)


def test_pdf_reads_page_by_page_joining_wrapped_lines(files):
    from lattice.documents import extract_pages, split_passages
    pages = extract_pages(files / "manual.pdf")
    assert [p for p, _ in pages] == [1, 2]
    p1 = split_passages(pages[0][1])
    assert "Use a damp cloth and never use solvents on the casing." in p1
    assert not any(p.startswith("Page 1") for p in p1)
    p2 = split_passages(pages[1][1])
    assert p2 == ["The warranty on the Vexlo heater lasts three years from "
                  "the date of purchase. Keep your receipt as proof."]


def test_hyphenated_line_ends_and_furniture(files):
    from lattice.documents import split_passages
    text = ("Chapter 3\nThe committee reviewed all of the infor-\nmation "
            "before voting on the plan.\n\n12  34  56\n...........\n")
    assert split_passages(text) == [
        "The committee reviewed all of the information before voting on "
        "the plan."]


def test_unreadable_files_fail_with_a_reason(tmp_path):
    from lattice.documents import DocumentError, extract_pages
    with pytest.raises(DocumentError, match="no file at"):
        extract_pages(tmp_path / "missing.pdf")
    bad = tmp_path / "sheet.xlsx"
    bad.write_bytes(b"x")
    with pytest.raises(DocumentError, match="can't read .xlsx"):
        extract_pages(bad)
    fake = tmp_path / "fake.pdf"
    fake.write_bytes(b"not a pdf at all")
    with pytest.raises(DocumentError, match="doesn't open as a PDF"):
        extract_pages(fake)


# ─── memory ─────────────────────────────────────────────────────────

def test_reading_stores_passages_with_file_and_page(telp, files):
    from lattice.documents import documents, read_document
    r = read_document(telp.agent, files / "manual.pdf")
    assert r["error"] is None and r["pages"] == 2 and r["added"] == 3
    lat = telp.agent.lattice
    rows = [(s, lat.get(i)) for i, s in zip(lat._ids, lat._sources)
            if s == "doc:manual.pdf"]
    assert len(rows) == 3
    assert {row["metadata"]["page"] for _, row in rows} == {1, 2}
    assert documents(telp.agent) == [{"name": "manual.pdf", "passages": 3}]


def test_rereading_replaces_not_duplicates(telp, files):
    from lattice.documents import documents, read_document
    read_document(telp.agent, files / "notes.txt")
    (files / "notes.txt").write_text("Rehearsal moved to Friday at the new "
                                     "hall by the river.\n")
    r = read_document(telp.agent, files / "notes.txt")
    assert r["replaced"] == 2 and r["added"] == 1
    assert documents(telp.agent) == [{"name": "notes.txt", "passages": 1}]
    assert not any("Thursday" in t for t in telp.agent.lattice._texts)


def test_forgetting_a_file_erases_only_its_passages(telp, files):
    from lattice.documents import (documents, find_document, forget_document,
                                   read_document)
    read_document(telp.agent, files / "report.docx")
    read_document(telp.agent, files / "manual.pdf")
    assert find_document(telp.agent, "report") == "report.docx"
    assert find_document(telp.agent, "REPORT.DOCX") == "report.docx"
    assert find_document(telp.agent, "budget") is None
    gone = forget_document(telp.agent, "report.docx")
    assert len(gone) == 4
    assert [d["name"] for d in documents(telp.agent)] == ["manual.pdf"]


# ─── the conversation ───────────────────────────────────────────────

def _last(t):
    return t.agent.turns[-1] if t.agent.turns else {}


def test_read_in_chat_then_answer_from_the_file_with_citation(telp, files):
    reply = telp.respond(f"read {files / 'report.docx'}")
    assert "report.docx" in reply and "4 passages" in reply
    reply = telp.respond("who led the Harlow bridge project?")
    assert "Dana Okafor" in reply
    reply = telp.respond("how do you know that?")
    assert _last(telp)["domain"] == "provenance"
    assert "the file report.docx" in reply


def test_pdf_answer_names_the_page(telp, files):
    telp.respond(f"read {files / 'manual.pdf'}")
    reply = telp.respond("how long is the warranty on the Vexlo heater?")
    assert "three years" in reply
    assert _last(telp)["domain"] == "document"
    assert "manual.pdf (page 2)" in reply


def test_bare_path_and_windows_style_command_both_read(telp, files):
    reply = telp.respond(str(files / "notes.txt"))
    assert "notes.txt" in reply and "2 passages" in reply
    reply = telp.respond(f"open the file '{files / 'manual.pdf'}'")
    assert "manual.pdf" in reply


def test_missing_and_unreadable_files_are_reported_not_remembered(telp,
                                                                    tmp_path):
    reply = telp.respond(f"read {tmp_path / 'nothing.pdf'}")
    assert "can't find a file" in reply
    assert _last(telp)["domain"] == "abstain"
    bad = tmp_path / "scan.pdf"
    bad.write_bytes(b"%PDF-1.4 garbage")
    reply = telp.respond(f"read {bad}")
    assert "couldn't read scan.pdf" in reply
    assert not any(s.startswith("doc:") for s in telp.agent.lattice._sources)


def test_list_and_forget_files_in_chat(telp, files):
    assert "haven't read any files" in telp.respond("what files have you read?")
    telp.respond(f"read {files / 'report.docx'}")
    telp.respond(f"read {files / 'manual.pdf'}")
    listing = telp.respond("which documents have you read?")
    assert "report.docx (4 passages)" in listing and "manual.pdf" in listing
    reply = telp.respond("forget the file report.docx")
    assert "forgotten the file 'report.docx'" in reply
    assert "report.docx" not in telp.respond("what files have you read?")
    reply = telp.respond("who led the Harlow bridge project?")
    assert "Dana Okafor" not in reply


def test_unrelated_questions_do_not_quote_the_file(telp, files):
    telp.respond(f"read {files / 'manual.pdf'}")
    reply = telp.respond("what is the capital of France?")
    assert "Vexlo" not in reply and "manual.pdf" not in reply


def test_a_question_mentioning_a_pdf_is_not_a_read_command(telp, files):
    telp.respond(f"read {files / 'manual.pdf'}")
    reply = telp.respond("what does manual.pdf say about the warranty?")
    assert "three years" in reply


# ─── the front door ─────────────────────────────────────────────────

def test_cli_read_documents_ask_and_forget(tmp_path, files):
    state = tmp_path / "state"
    r = _telp(state, "read", str(files / "report.docx"), str(files / "manual.pdf"))
    assert r.returncode == 0, r.stderr
    assert "report.docx: 4 passages remembered" in r.stdout
    assert "manual.pdf (2 pages): 3 passages remembered" in r.stdout
    r = _telp(state, "documents")
    assert "report.docx: 4 passages" in r.stdout
    r = _telp(state, "ask", "who led the Harlow bridge project?")
    assert "Dana Okafor" in r.stdout
    r = _telp(state, "read", str(files / "missing.pdf"))
    assert r.returncode == 1 and "no file at" in r.stdout
